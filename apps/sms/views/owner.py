from uuid import uuid4

from django.conf import settings
from django.db import transaction
from django.db.models import Q
from django.shortcuts import get_object_or_404
from rest_framework import parsers, permissions, serializers, status, views
from rest_framework.response import Response
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import extend_schema, inline_serializer

from apps.market.access import accessible_markets
from apps.sms.models import Contact, Line, SmsCampaign, SmsRecipient, Template
from apps.sms.serializers.owner import (
    CampaignCreateSerializer, SmsContactSerializer, LineListSerializer,
    SmsCampaignSerializer, SmsTariffSerializer, TemplateCreateSerializer,
    TemplateListSerializer, normalize_mobile,
)
from apps.sms.services import active_tariff, estimate, reserve_campaign, sms_eligibility
from utils.response import ApiResponse


def api_data(data, code=200):
    return Response(ApiResponse(success=True, code=code, data=data), status=code)


class SmsDashboardView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: OpenApiTypes.OBJECT}, tags=['SMS Panel'])
    def get(self, request):
        campaigns = SmsCampaign.objects.filter(user=request.user)
        tariff = active_tariff()
        return api_data({
            'first_use': not campaigns.exists(),
            'eligibility': sms_eligibility(request.user),
            'campaign_count': campaigns.count(),
            'tariff': SmsTariffSerializer(tariff).data if tariff else None,
        })


class LineListView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: LineListSerializer(many=True)}, tags=['SMS Panel'])
    def get(self, request):
        return api_data(LineListSerializer(Line.objects.filter(is_active=True), many=True).data)


class TemplateListView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: TemplateListSerializer(many=True)}, tags=['SMS Panel'])
    def get(self, request):
        templates = Template.objects.filter(is_active=True).filter(
            Q(is_public=True, approval_status='approved') | Q(owner=request.user)
        ).order_by('title')
        return api_data(TemplateListSerializer(templates, many=True).data)

    @extend_schema(
        request=TemplateCreateSerializer,
        responses={201: TemplateListSerializer},
        tags=['SMS Panel'],
    )
    def post(self, request):
        serializer = TemplateCreateSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        template = serializer.save()
        return api_data(TemplateListSerializer(template).data, status.HTTP_201_CREATED)


class ContactListCreateView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: SmsContactSerializer(many=True)}, tags=['SMS Panel'])
    def get(self, request):
        contacts = Contact.objects.filter(user=request.user).order_by('name', 'mobile_number')
        return api_data(SmsContactSerializer(contacts, many=True).data)

    @extend_schema(
        request=SmsContactSerializer,
        responses={201: SmsContactSerializer},
        tags=['SMS Panel'],
    )
    def post(self, request):
        serializer = SmsContactSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        contact, _ = Contact.objects.update_or_create(
            user=request.user,
            mobile_number=serializer.validated_data['mobile_number'],
            defaults={
                'name': serializer.validated_data.get('name'),
                'source': serializer.validated_data.get('source', 'phone'),
                'has_consent': serializer.validated_data.get('has_consent', False),
            },
        )
        return api_data(SmsContactSerializer(contact).data, status.HTTP_201_CREATED)


class ContactExcelUploadView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)
    parser_classes = (parsers.MultiPartParser,)

    @extend_schema(
        request=inline_serializer(
            name='SmsContactExcelUpload',
            fields={'file': serializers.FileField()},
        ),
        responses={200: OpenApiTypes.OBJECT},
        tags=['SMS Panel'],
    )
    def post(self, request):
        upload = request.FILES.get('file')
        if not upload or not upload.name.lower().endswith('.xlsx'):
            raise serializers.ValidationError({'file': 'فایل با پسوند xlsx لازم است.'})
        if upload.size > 5 * 1024 * 1024:
            raise serializers.ValidationError({'file': 'حجم فایل نباید بیشتر از ۵ مگابایت باشد.'})
        try:
            from openpyxl import load_workbook
            sheet = load_workbook(upload, read_only=True, data_only=True).active
        except Exception as exc:
            raise serializers.ValidationError({'file': 'فایل اکسل معتبر نیست.'}) from exc
        imported, rejected = [], []
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        with transaction.atomic():
            for index, row in enumerate(rows, start=2):
                name = str(row[0] or '').strip() if row else ''
                raw_mobile = row[1] if row and len(row) > 1 else None
                try:
                    mobile = normalize_mobile(raw_mobile)
                except serializers.ValidationError:
                    rejected.append({'row': index, 'mobile_number': str(raw_mobile or '')})
                    continue
                contact, _ = Contact.objects.update_or_create(
                    user=request.user, mobile_number=mobile,
                    defaults={'name': name, 'source': 'excel', 'has_consent': True},
                )
                imported.append(contact)
        return api_data({
            'imported_count': len(imported), 'rejected': rejected,
            'contacts': SmsContactSerializer(imported, many=True).data,
        })


class CampaignListCreateView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: SmsCampaignSerializer(many=True)}, tags=['SMS Panel'])
    def get(self, request):
        campaigns = SmsCampaign.objects.filter(user=request.user).prefetch_related('recipients')
        return api_data(SmsCampaignSerializer(campaigns, many=True).data)

    @extend_schema(
        request=CampaignCreateSerializer,
        responses={201: SmsCampaignSerializer},
        tags=['SMS Panel'],
    )
    @transaction.atomic
    def post(self, request):
        if getattr(settings, 'SMS_BILLING_ENABLED', True) is False:
            return Response(ApiResponse(success=False, code=503, error={
                'code': 'sms_billing_unavailable', 'detail': 'SMS billing is disabled.',
            }), status=status.HTTP_503_SERVICE_UNAVAILABLE)
        eligibility = sms_eligibility(request.user)
        if not eligibility['eligible']:
            return Response(ApiResponse(success=False, code=403, error={
                'code': 'sms_identity_not_approved', 'detail': eligibility,
            }), status=status.HTTP_403_FORBIDDEN)
        serializer = CampaignCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        tariff = active_tariff()
        if not tariff:
            return Response(ApiResponse(success=False, code=503, error={
                'code': 'sms_tariff_unavailable', 'detail': 'تعرفه فعال ثبت نشده است.',
            }), status=status.HTTP_503_SERVICE_UNAVAILABLE)
        line = get_object_or_404(Line.objects.filter(is_active=True), id=data['line'])
        market = None
        if data.get('market'):
            market = get_object_or_404(accessible_markets(request.user, write=True), id=data['market'])
        template = None
        campaign_status = SmsCampaign.WAITING_APPROVAL
        if data.get('template'):
            template = get_object_or_404(Template, id=data['template'], is_active=True)
            if not (template.owner_id == request.user.id or template.is_public):
                return Response(status=status.HTTP_403_FORBIDDEN)
            if template.approval_status != 'approved':
                raise serializers.ValidationError({'template': 'متن هنوز تأیید نشده است.'})
            campaign_status = SmsCampaign.READY_FOR_PAYMENT
        else:
            template = Template.objects.create(
                owner=request.user, title=data['title'], category='campaign',
                content=data['message'], variables=[], estimated_cost=0,
                approval_status='pending', is_public=False,
            )
        _, segments, cost = estimate(data['message'], data.get('link', ''), len(data['recipients']), tariff)
        key = request.headers.get('Idempotency-Key') or uuid4().hex
        existing = SmsCampaign.objects.filter(user=request.user, idempotency_key=key).first()
        if existing:
            return api_data(SmsCampaignSerializer(existing).data)
        campaign = SmsCampaign.objects.create(
            user=request.user, market=market, template=template, tariff=tariff, line=line,
            title=data['title'], message=data['message'], link=data.get('link', ''),
            status=campaign_status, recipient_count=len(data['recipients']),
            segment_count=segments, estimated_cost=cost, idempotency_key=key,
            scheduled_at=data.get('scheduled_at'),
        )
        SmsRecipient.objects.bulk_create([
            SmsRecipient(campaign=campaign, **recipient) for recipient in data['recipients']
        ])
        return api_data(SmsCampaignSerializer(campaign).data, status.HTTP_201_CREATED)


class CampaignDetailView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(responses={200: SmsCampaignSerializer}, tags=['SMS Panel'])
    def get(self, request, pk):
        campaign = get_object_or_404(
            SmsCampaign.objects.prefetch_related('recipients'), id=pk, user=request.user,
        )
        return api_data(SmsCampaignSerializer(campaign).data)


class CampaignPayView(views.APIView):
    permission_classes = (permissions.IsAuthenticated,)

    @extend_schema(request=None, responses={200: SmsCampaignSerializer}, tags=['SMS Panel'])
    def post(self, request, pk):
        campaign = get_object_or_404(SmsCampaign, id=pk, user=request.user)
        campaign, paid, error = reserve_campaign(campaign, request.user)
        if not paid:
            return Response(ApiResponse(success=False, code=409, error={
                'code': error, 'detail': 'پرداخت یا رزرو هزینه انجام نشد.',
            }), status=status.HTTP_409_CONFLICT)
        return api_data(SmsCampaignSerializer(campaign).data)


# Compatibility endpoints now use the safe campaign workflow instead of direct sending.
BulkSmsView = CampaignListCreateView
PatternSmsView = CampaignListCreateView
